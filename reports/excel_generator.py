"""
APEX LAB Attendance System — Weekly Excel Report Generator
Generates professional .xlsx reports with Summary, Daily Log,
Unknown Faces, and Analytics sheets.
"""

import logging
from datetime import datetime, timedelta
from collections import defaultdict, Counter
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side,
)
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import config

from database.models import AttendanceRecord, Person

logger = logging.getLogger(__name__)

# ── Reusable styles ──────────────────────────────────────────────────────────

_HEADER_FILL = PatternFill(
    start_color=config.BG_COLOR_CARD.lstrip("#"),
    end_color=config.BG_COLOR_CARD.lstrip("#"),
    fill_type="solid",
)
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

_GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _parse_timestamp(ts: str) -> Tuple[str, str]:
    """Split an ISO-8601 timestamp into (date_str 'YYYY-MM-DD', time_str 'HH:MM:SS').

    Handles both 'YYYY-MM-DDTHH:MM:SS' and 'YYYY-MM-DD HH:MM:SS' formats.
    """
    ts = ts.replace("T", " ")
    parts = ts.split(" ", 1)
    date_str = parts[0]
    time_str = parts[1] if len(parts) > 1 else "00:00:00"
    return date_str, time_str


class WeeklyReportGenerator:
    """Generates a weekly attendance report as an Excel workbook."""

    DAY_NAMES = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

    def __init__(self, db_manager):
        """
        Parameters
        ----------
        db_manager : database.db_manager.DatabaseManager
            Provides query methods for attendance data.
        """
        self.db = db_manager

    # ── Public API ────────────────────────────────────────────────────────

    def generate_report(self, start_date: str, end_date: str) -> str:
        """
        Build the full weekly report and save it to *config.REPORTS_DIR*.

        Parameters
        ----------
        start_date : str  – "YYYY-MM-DD"
        end_date   : str  – "YYYY-MM-DD"

        Returns
        -------
        str – Absolute path of the generated .xlsx file.
        """
        start_dt = datetime.strptime(start_date, "%Y-%m-%d").date()
        end_dt = datetime.strptime(end_date, "%Y-%m-%d").date()

        wb = Workbook()

        # Fetch data once
        persons: List[Person] = self.db.get_all_persons()
        records: List[AttendanceRecord] = self.db.get_attendance_for_week(start_date, end_date)
        unknown_records: List[AttendanceRecord] = self.db.get_unknown_entries(start_date, end_date)

        # Build a name lookup for person_id → name
        person_map: Dict[int, Person] = {p.id: p for p in persons}

        self._build_summary_sheet(wb, persons, records, start_dt, end_dt)
        self._build_daily_log_sheet(wb, records, person_map)
        self._build_unknown_faces_sheet(wb, unknown_records)
        self._build_analytics_sheet(wb, persons, records, person_map, start_dt, end_dt)

        # Save
        filename = config.REPORT_FILENAME_FORMAT.format(
            start=start_date, end=end_date,
        )
        filepath = config.REPORTS_DIR / filename
        filepath.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(filepath))
        logger.info("Report saved → %s", filepath)
        return str(filepath)

    # ── Sheet builders ────────────────────────────────────────────────────

    def _build_summary_sheet(self, wb, persons: List[Person],
                             records: List[AttendanceRecord],
                             start_dt, end_dt):
        """Sheet 1 — per-person weekly summary grid."""
        ws = wb.active
        ws.title = "Summary"

        # Header row
        headers = ["Name", "Role"] + self.DAY_NAMES + ["Total", "%"]
        ws.append(headers)
        self._style_header(ws, 1)

        # Build lookup: (person_id, date_str) → earliest check-in time_str
        checkin_map: Dict[Tuple[int, str], str] = {}
        for rec in records:
            if rec.person_id is None or rec.status == "unknown":
                continue
            date_str, time_str = _parse_timestamp(rec.timestamp)
            key = (rec.person_id, date_str)
            if key not in checkin_map or time_str < checkin_map[key]:
                checkin_map[key] = time_str

        # Date list for the week
        date_range = [
            start_dt + timedelta(days=i)
            for i in range((end_dt - start_dt).days + 1)
        ]

        for person in persons:
            pid = person.id
            row: list = [person.name, person.role]
            present_count = 0
            work_day_count = 0

            for d in date_range:
                weekday = d.weekday()  # 0=Mon
                day_str = d.strftime("%Y-%m-%d")
                is_work_day = weekday in config.WORK_DAYS

                if not is_work_day:
                    row.append("—")
                    continue

                work_day_count += 1
                key = (pid, day_str)
                if key in checkin_map:
                    present_count += 1
                    t = checkin_map[key][:5]  # "HH:MM"
                    row.append(f"✓ {t}")
                else:
                    row.append("✗")

            row.append(present_count)
            pct = (present_count / work_day_count * 100) if work_day_count else 0
            row.append(round(pct, 1))
            ws.append(row)

        # Percentage column formatting
        pct_col = len(headers)
        self._add_conditional_formatting(ws, pct_col, len(persons))

        self._auto_width(ws)
        ws.freeze_panes = "C2"

    def _build_daily_log_sheet(self, wb, records: List[AttendanceRecord],
                               person_map: Dict[int, Person]):
        """Sheet 2 — every attendance event sorted by date / time."""
        ws = wb.create_sheet("Daily Log")
        headers = ["Date", "Time", "Name", "Status", "Confidence"]
        ws.append(headers)
        self._style_header(ws, 1)

        # Sort by timestamp
        sorted_recs = sorted(records, key=lambda r: r.timestamp)
        for rec in sorted_recs:
            date_str, time_str = _parse_timestamp(rec.timestamp)
            name = "Unknown"
            if rec.person_id is not None and rec.person_id in person_map:
                name = person_map[rec.person_id].name
            ws.append([
                date_str,
                time_str[:5],
                name,
                rec.status,
                round(rec.confidence, 2) if rec.confidence else "",
            ])

        self._auto_width(ws)
        ws.freeze_panes = "A2"

    def _build_unknown_faces_sheet(self, wb, unknown_records: List[AttendanceRecord]):
        """Sheet 3 — unknown / unrecognised face events."""
        ws = wb.create_sheet("Unknown Faces")
        headers = ["Date", "Time", "Snapshot Path", "Notes"]
        ws.append(headers)
        self._style_header(ws, 1)

        for rec in unknown_records:
            date_str, time_str = _parse_timestamp(rec.timestamp)
            ws.append([
                date_str,
                time_str[:5],
                rec.snapshot_path or "",
                f"Confidence: {rec.confidence:.2f}" if rec.confidence else "",
            ])

        self._auto_width(ws)
        ws.freeze_panes = "A2"

    def _build_analytics_sheet(self, wb, persons: List[Person],
                               records: List[AttendanceRecord],
                               person_map: Dict[int, Person],
                               start_dt, end_dt):
        """Sheet 4 — peak hours, average times, late arrivals, absences."""
        ws = wb.create_sheet("Analytics")

        # Build earliest-checkin map: (person_id, date_str) → time_str
        earliest: Dict[Tuple[int, str], str] = {}
        for rec in records:
            if rec.person_id is None or rec.status == "unknown":
                continue
            date_str, time_str = _parse_timestamp(rec.timestamp)
            key = (rec.person_id, date_str)
            if key not in earliest or time_str < earliest[key]:
                earliest[key] = time_str

        # ---- 1. Peak Check-In Hour ----
        ws.append(["Peak Check-In Hour"])
        self._style_header(ws, ws.max_row)

        hour_counter: Counter = Counter()
        for rec in records:
            if rec.status == "unknown":
                continue
            try:
                _, time_str = _parse_timestamp(rec.timestamp)
                hour = int(time_str.split(":")[0])
                hour_counter[hour] += 1
            except (ValueError, IndexError):
                continue

        if hour_counter:
            ws.append(["Hour", "Count"])
            for hour in sorted(hour_counter):
                label = f"{hour:02d}:00 – {hour:02d}:59"
                ws.append([label, hour_counter[hour]])
            peak = hour_counter.most_common(1)[0]
            ws.append([])
            ws.append([f"Peak hour: {peak[0]:02d}:00 ({peak[1]} check-ins)"])
        else:
            ws.append(["No check-in data available."])

        ws.append([])

        # ---- 2. Average Check-In Time Per Person ----
        ws.append(["Average Check-In Time Per Person"])
        self._style_header(ws, ws.max_row)
        ws.append(["Name", "Avg Check-In"])

        person_times: Dict[int, List[int]] = defaultdict(list)
        for (pid, _), tstr in earliest.items():
            parts = tstr.split(":")
            total_minutes = int(parts[0]) * 60 + int(parts[1])
            person_times[pid].append(total_minutes)

        for pid in sorted(person_times):
            minutes_list = person_times[pid]
            avg_min = sum(minutes_list) / len(minutes_list)
            h, m = divmod(int(avg_min), 60)
            name = person_map[pid].name if pid in person_map else f"ID {pid}"
            ws.append([name, f"{h:02d}:{m:02d}"])

        ws.append([])

        # ---- 3. Late Arrivals ----
        ws.append(["Late Arrivals"])
        self._style_header(ws, ws.max_row)
        ws.append(["Name", "Date", "Arrived", "Expected"])

        expected_parts = config.EXPECTED_CHECKIN_TIME.split(":")
        expected_minutes = int(expected_parts[0]) * 60 + int(expected_parts[1])

        late_count = 0
        for (pid, date_str), tstr in sorted(earliest.items(), key=lambda x: (x[0][1], x[0][0])):
            parts = tstr.split(":")
            arrived_minutes = int(parts[0]) * 60 + int(parts[1])
            if arrived_minutes > expected_minutes:
                late_count += 1
                name = person_map[pid].name if pid in person_map else f"ID {pid}"
                ws.append([name, date_str, tstr[:5], config.EXPECTED_CHECKIN_TIME])

        if late_count == 0:
            ws.append(["No late arrivals recorded."])

        ws.append([])

        # ---- 4. Consecutive Absences (≥ ABSENT_FLAG_DAYS days) ----
        ws.append([f"Absent {config.ABSENT_FLAG_DAYS}+ Consecutive Work Days"])
        self._style_header(ws, ws.max_row)
        ws.append(["Name", "From", "To", "Days"])

        work_dates = []
        d = start_dt
        while d <= end_dt:
            if d.weekday() in config.WORK_DAYS:
                work_dates.append(d)
            d += timedelta(days=1)

        flagged = False
        for person in persons:
            pid = person.id
            consecutive = 0
            streak_start = None

            for wd in work_dates:
                key = (pid, wd.strftime("%Y-%m-%d"))
                if key not in earliest:
                    if consecutive == 0:
                        streak_start = wd
                    consecutive += 1
                else:
                    if consecutive >= config.ABSENT_FLAG_DAYS:
                        streak_end = wd - timedelta(days=1)
                        ws.append([
                            person.name,
                            streak_start.strftime("%Y-%m-%d"),
                            streak_end.strftime("%Y-%m-%d"),
                            consecutive,
                        ])
                        flagged = True
                    consecutive = 0
                    streak_start = None

            # Handle streak extending to end of range
            if consecutive >= config.ABSENT_FLAG_DAYS and streak_start and work_dates:
                ws.append([
                    person.name,
                    streak_start.strftime("%Y-%m-%d"),
                    work_dates[-1].strftime("%Y-%m-%d"),
                    consecutive,
                ])
                flagged = True

        if not flagged:
            ws.append(["No persons flagged."])

        self._auto_width(ws)
        ws.freeze_panes = "A2"

    # ── Styling helpers ───────────────────────────────────────────────────

    @staticmethod
    def _style_header(ws, row: int):
        """Apply branded header styling to every cell in *row*."""
        for cell in ws[row]:
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGNMENT
            cell.border = _THIN_BORDER

    @staticmethod
    def _auto_width(ws):
        """Auto-adjust column widths based on the longest value."""
        for col_cells in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                try:
                    cell_len = len(str(cell.value)) if cell.value is not None else 0
                    if cell_len > max_len:
                        max_len = cell_len
                except Exception:
                    pass
                # Apply border to every cell while we iterate
                cell.border = _THIN_BORDER
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

    @staticmethod
    def _add_conditional_formatting(ws, pct_col: int, person_count: int):
        """
        Add green fill for 100 % attendance and red fill for < 60 %
        on the percentage column.
        """
        col_letter = get_column_letter(pct_col)
        cell_range = f"{col_letter}2:{col_letter}{person_count + 1}"

        ws.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="equal",
                formula=["100"],
                fill=_GREEN_FILL,
            ),
        )
        ws.conditional_formatting.add(
            cell_range,
            CellIsRule(
                operator="lessThan",
                formula=["60"],
                fill=_RED_FILL,
            ),
        )
